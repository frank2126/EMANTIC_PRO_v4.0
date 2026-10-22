# ═══════════════════════════════════════════════════════════
#  TESTS DPV — carga del Excel, listado y DPV mensual (endpoints)
#  Hoja correcta dentro del libro · limpieza de datos · km mensual · DPV calculado en vivo
#  (las reglas puras están en test_dpv_reglas.py)
#
#  Cada prueba usa ids y meses propios (2030, 2031, 2032…): no dependen del orden ni de que la BD se reinicie.
# ═══════════════════════════════════════════════════════════

import io
from datetime import datetime, time

from fastapi import status
from openpyxl import Workbook

from routes import datos


def doble_codificacion(texto: str) -> str:
    """Simula el daño del Excel de Transmitools: UTF-8 leído como latin-1."""
    return texto.encode("utf-8").decode("latin-1")


HEADERS = [
    "DPV", "ESTADO", "F. INICIO DPV", "F. CIERRE DPV", "EMPRESA", "FUENTE", "PLACA", "MOVIL", "TIPOLOGIA",
    "FECHA INMOVILIZACION", "HORA", "CAUSA DE INMOVILIZACION", "DESCRIPCION DE LA NOVEDAD", "RUTA", "OPERADOR",
    "DIA SEMANA", "ALERTA", "FRANJA HORARIA",
]


def fila(dpv, estado="No Contestado", fecha=datetime(2030, 1, 5), tipologia="PADRON", empresa="EMASIVO 10 S.A.S.",
         movil="Z32-9001", hora=time(22, 4), descripcion="Sin novedad", ruta="Ruta A100", operador=320000,
         alerta=None, franja="P.\u00a0M.", dia="LUNES"):
    return [
        dpv, estado, datetime(2031, 1, 7), datetime(2031, 1, 14, 23, 59), empresa, "URBANO", "AAA000", movil,
        tipologia, fecha, hora, "LIMPIAPARABRISAS", descripcion, ruta, operador, dia, alerta, franja,
    ]


def excel_dpv(filas, headers=HEADERS) -> bytes:
    """Libro como el real: la hoja ACTIVA es otra (km por día) y los eventos van en «base de datos»."""
    wb = Workbook()
    km = wb.active
    km.title = "km dia sca vw e10"
    km.append(["MOVIL ", "LINEA ", datetime(2031, 1, 1), datetime(2031, 1, 2)])
    km.append(["Z32-9001", "SCANIA ", 0, 231])
    base = wb.create_sheet("base de datos")
    base.append(headers)
    for f in filas:
        base.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


MENSUAL_HEADERS = ["MES", "TOTAL KM", "TOTAL INNMOVILIZADOS", "DPV GENERAL", "FECHA CORTE", "KM PADRON", "KM BUSETON",
                   "INM PADRON", "INM BUSETON", "DPV PADRON", "DPV BUSETON"]


def excel_mensual(e10, e16=()) -> bytes:
    wb = Workbook()
    for nombre, filas in (("DPV_E10", e10), ("DPV_E16", e16)):
        ws = wb.active if nombre == "DPV_E10" else wb.create_sheet()
        ws.title = nombre
        ws.append(MENSUAL_HEADERS)
        for f in filas:
            ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def subir(client, headers, url, contenido, nombre="DPV.xlsx"):
    return client.post(url, files={"file": (nombre, contenido, XLSX)}, headers=headers)


# ═══════════════════════════════════════════════════════════
#  CARGA DE EVENTOS
# ═══════════════════════════════════════════════════════════

class TestCargaDPV:

    def test_usa_la_hoja_de_eventos_aunque_no_sea_la_activa(self, client, admin_headers):
        """✅ El libro del DPV trae 17 hojas: se usa la que tiene DPV / ESTADO / PLACA"""
        r = subir(client, admin_headers, "/api/dpv/upload", excel_dpv([fila(7001), fila(7002)]))
        assert r.status_code == status.HTTP_200_OK, r.text
        assert r.json()["insertados"] == 2 and r.json()["errores"] == 0
        assert r.json()["hoja"] == "base de datos"

    def test_limpia_los_datos_al_cargar(self, client, admin_headers):
        """✅ Espacios, mayúsculas, franja con NBSP, hora con fecha, alerta con emoji, operador decimal y texto dañado"""
        sucio = fila(
            7101, estado="no  contestado ", tipologia="Buseton ", ruta="Ruta  A100 ",
            descripcion=doble_codificacion("Revisión  de la señal"), operador=320000.0,
            hora=datetime(1900, 1, 1, 7, 5, 9), franja="A.\u00a0M.", alerta="⚠️ Móvil inválido:operdor inexistente",
        )
        sin_franja = fila(7102, hora="7:05", franja=None, dia=None)   # sin franja ni día → se deducen de la hora y la fecha
        r = subir(client, admin_headers, "/api/dpv/upload", excel_dpv([sucio, sin_franja]))
        assert r.status_code == status.HTTP_200_OK, r.text
        assert r.json()["calidad"]["textos_reparados"] >= 1

        r = client.get("/api/dpv?slim=true&limit=100000", headers=admin_headers)
        items = {i["dpv_id"]: i for i in r.json()["items"]}
        item = items[7101]
        assert item["estado"] == "No Contestado"
        assert item["tipologia"] == "BUSETON"
        assert item["ruta"] == "Ruta A100"
        assert item["descripcion_novedad"] == "Revisión de la señal"
        assert item["operador"] == "320000"
        assert item["hora"] == "07:05:09" and item["franja_horaria"] == "AM"
        assert item["alerta"] == "Móvil inválido: operador inexistente"
        assert items[7102]["hora"] == "07:05:00" and items[7102]["franja_horaria"] == "AM"
        assert items[7102]["dia_semana"] == "SÁBADO"                 # 5-ene-2030 cae sábado

    def test_el_listado_trae_lo_que_usa_el_dashboard(self, client, admin_headers):
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([fila(7201)]))
        item = {i["dpv_id"]: i for i in client.get("/api/dpv?slim=true&limit=100000", headers=admin_headers)
                .json()["items"]}[7201]
        assert item["fecha_inmovilizacion"].startswith("2030-01-05")
        assert item["fecha_cierre"].startswith("2031-01-14")
        assert item["dia_semana"] == "LUNES" and item["franja_horaria"] == "PM"
        assert item["cuenta"] is True                                   # No Contestado cuenta como inmovilización

    def test_marca_si_el_evento_cuenta_para_el_dpv(self, client, admin_headers):
        """✅ El backend dice qué eventos cuentan (todos menos «Contestado Contundente»), también en el listado completo"""
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([
            fila(7211, estado="Contestado Contundente"), fila(7212, estado="Validado"), fila(7213),
        ]))
        for url in ("/api/dpv?slim=true&limit=100000", "/api/dpv?limit=100000"):
            items = {i["dpv_id"]: i for i in client.get(url, headers=admin_headers).json()["items"]}
            assert (items[7211]["cuenta"], items[7212]["cuenta"], items[7213]["cuenta"]) == (False, True, True), url

    def test_recarga_actualiza_sin_duplicar(self, client, admin_headers):
        """✅ Subir el mismo Excel dos veces actualiza por DPV (upsert)"""
        contenido = excel_dpv([fila(7301), fila(7302)])
        primera = subir(client, admin_headers, "/api/dpv/upload", contenido).json()
        segunda = subir(client, admin_headers, "/api/dpv/upload", contenido).json()
        assert primera["insertados"] == 2
        assert segunda["insertados"] == 0 and segunda["actualizados"] == 2

    def test_sin_las_columnas_requeridas_el_error_es_claro(self, client, admin_headers):
        """❌ Un libro sin hoja de eventos da 400 y dice qué falta"""
        r = subir(client, admin_headers, "/api/dpv/upload", excel_dpv([], headers=["A", "B", "C"]))
        assert r.status_code == status.HTTP_400_BAD_REQUEST
        assert "Faltan columnas requeridas" in r.json()["detail"]

    def test_carga_requiere_admin(self, client, tecnico_headers):
        """❌ Un técnico no puede cargar el Excel"""
        r = subir(client, tecnico_headers, "/api/dpv/upload", excel_dpv([fila(7401)]))
        assert r.status_code == status.HTTP_403_FORBIDDEN

    def test_listado_requiere_autenticacion(self, client):
        """❌ GET /api/dpv sin token falla"""
        assert client.get("/api/dpv").status_code == status.HTTP_403_FORBIDDEN

    def test_listado_por_mes(self, client, admin_headers):
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([
            fila(7501, fecha=datetime(2033, 4, 30, 23, 0)), fila(7502, fecha=datetime(2033, 5, 1)),
        ]))
        abril = client.get("/api/dpv?mes=2033-04&slim=true&limit=100", headers=admin_headers).json()
        assert [i["dpv_id"] for i in abril["items"]] == [7501]


# ═══════════════════════════════════════════════════════════
#  KM MENSUAL Y DPV EN VIVO
# ═══════════════════════════════════════════════════════════

def serie(client, headers, unidad="e10"):
    r = client.get("/api/dpv-mensual", headers=headers)
    assert r.status_code == status.HTTP_200_OK, r.text
    return {m["mes"]: m for m in r.json()[unidad]}


class TestDpvMensual:

    def test_carga_los_km_del_excel_y_omite_los_meses_sin_km(self, client, admin_headers):
        """✅ Un mes sin TOTAL KM no se guarda como 0: simplemente no hay km"""
        contenido = excel_mensual([
            [datetime(2034, 1, 1), 400000, 4, 100000, datetime(2034, 1, 31), 300000, 100000, 3, 1, 100000, 100000],
            [datetime(2034, 2, 1), None, 5, None, None, None, None, None, None, None, None],
        ])
        r = subir(client, admin_headers, "/api/dpv-mensual/upload", contenido)
        assert r.status_code == status.HTTP_200_OK, r.text
        d = r.json()
        assert d["formato"] == "nuevo" and d["insertados"] == 1 and d["errores"] == 0
        assert d["sin_km"]["E10"] == ["2034-02"]
        assert "2034-02" not in serie(client, admin_headers)           # sin km y sin eventos → no hay fila

    def test_el_dpv_se_calcula_en_vivo_desde_los_eventos(self, client, admin_headers):
        """✅ km ÷ inmovilizaciones; «Contestado Contundente» no cuenta; el último día del mes con hora sí"""
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([
            fila(8001, fecha=datetime(2031, 1, 5)),
            fila(8002, fecha=datetime(2031, 1, 20), tipologia="PADRON"),
            fila(8003, fecha=datetime(2031, 1, 31, 18, 45), tipologia="PADRON"),        # último día, con hora
            fila(8004, fecha=datetime(2031, 1, 10), tipologia="BUSETON ", estado="Validado"),
            fila(8005, fecha=datetime(2031, 1, 11), estado="Contestado Contundente"),   # no cuenta
            fila(8006, fecha=datetime(2031, 2, 1)),                                     # otro mes
        ]))
        subir(client, admin_headers, "/api/dpv-mensual/upload", excel_mensual([
            [datetime(2031, 1, 1), 400000, 99, 1, None, 300000, 100000, 9, 9, 1, 1],    # lo que dice el Excel se ignora
        ]))
        ene = serie(client, admin_headers)["2031-01"]
        assert (ene["eventos"], ene["descartadas"], ene["inm"]) == (5, 1, 4)
        assert (ene["inm_padron"], ene["inm_buseton"]) == (3, 1)
        assert ene["dpv"] == 100000.0                      # 400.000 / 4
        assert ene["dpv_padron"] == 100000.0               # 300.000 / 3
        assert ene["dpv_buseton"] == 100000.0              # 100.000 / 1
        assert ene["nivel"] == "cumple" and ene["origen"] == "excel"
        assert ene["inm_excel"] == 99                      # solo como referencia

        feb = serie(client, admin_headers)["2031-02"]      # hay eventos pero no km → vacío, no 0
        assert feb["inm"] == 1 and feb["dpv"] is None and feb["origen"] == "sin_km"

    def test_un_evento_que_pasa_a_contundente_mejora_el_dpv(self, client, admin_headers):
        """✅ Recargar el Excel con el estado cambiado actualiza el DPV del mes sin tocar los km"""
        subir(client, admin_headers, "/api/dpv/upload",
              excel_dpv([fila(8101, fecha=datetime(2032, 3, 2)), fila(8102, fecha=datetime(2032, 3, 3))]))
        subir(client, admin_headers, "/api/dpv-mensual/upload",
              excel_mensual([[datetime(2032, 3, 1), 20000, 2, 10000, None, None, None, None, None, None, None]]))
        assert serie(client, admin_headers)["2032-03"]["dpv"] == 10000.0
        subir(client, admin_headers, "/api/dpv/upload",
              excel_dpv([fila(8101, fecha=datetime(2032, 3, 2)),
                         fila(8102, estado="Contestado Contundente", fecha=datetime(2032, 3, 3))]))
        assert serie(client, admin_headers)["2032-03"]["dpv"] == 20000.0

    def test_la_serie_de_cada_unidad_es_independiente(self, client, admin_headers):
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([
            fila(8201, fecha=datetime(2035, 6, 2)),
            fila(8202, fecha=datetime(2035, 6, 3), empresa="EMASIVO 16 S.A.S.", movil="Z34-9001"),
            fila(8203, fecha=datetime(2035, 6, 4), empresa="EMASIVO 16 S.A.S.", movil="Z34-9001"),
        ]))
        subir(client, admin_headers, "/api/dpv-mensual/upload", excel_mensual(
            [[datetime(2035, 6, 1), 10000, 1, 10000, None, None, None, None, None, None, None]],
            [[datetime(2035, 6, 1), 30000, 2, 15000, None, None, None, None, None, None, None]]))
        assert serie(client, admin_headers, "e10")["2035-06"]["dpv"] == 10000.0
        assert serie(client, admin_headers, "e16")["2035-06"]["dpv"] == 15000.0

    def test_carga_mensual_requiere_admin(self, client, tecnico_headers):
        r = subir(client, tecnico_headers, "/api/dpv-mensual/upload", excel_mensual([]))
        assert r.status_code == status.HTTP_403_FORBIDDEN

    def test_responde_los_umbrales(self, client, admin_headers):
        d = client.get("/api/dpv-mensual", headers=admin_headers).json()
        assert d["umbrales"] == {"estandar": 10000, "critico": 7000}
        assert d["estados_descartados"] == ["Contestado Contundente"]


class TestCapturaManualDeKm:

    def test_guarda_y_corrige_el_km_de_un_mes(self, client, admin_headers):
        """✅ Para los meses que el Excel aún no trae"""
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([
            fila(8301, fecha=datetime(2036, 2, 2)), fila(8302, fecha=datetime(2036, 2, 3)),
        ]))
        r = client.put("/api/dpv-km", json={"empresa": "e10", "mes": "2036-02", "km_total": 30000}, headers=admin_headers)
        assert r.status_code == status.HTTP_200_OK, r.text
        assert r.json() == {"empresa": "E10", "mes": "2036-02", "actualizado": False}
        feb = serie(client, admin_headers)["2036-02"]
        assert feb["dpv"] == 15000.0 and feb["origen"] == "manual"

        r = client.put("/api/dpv-km", json={"empresa": "E10", "mes": "2036-02", "km_total": 40000}, headers=admin_headers)
        assert r.json()["actualizado"] is True
        assert serie(client, admin_headers)["2036-02"]["dpv"] == 20000.0

    def test_borrar_devuelve_el_mes_a_sin_km(self, client, admin_headers):
        subir(client, admin_headers, "/api/dpv/upload", excel_dpv([fila(8401, fecha=datetime(2037, 5, 2))]))
        client.put("/api/dpv-km", json={"empresa": "E10", "mes": "2037-05", "km_total": 9000}, headers=admin_headers)
        assert client.delete("/api/dpv-km/E10/2037-05", headers=admin_headers).json() == {"eliminado": True}
        assert serie(client, admin_headers)["2037-05"]["dpv"] is None
        assert client.delete("/api/dpv-km/E10/2037-05", headers=admin_headers).json() == {"eliminado": False}

    def test_valida_unidad_mes_y_kilometros(self, client, admin_headers):
        """❌ Datos que no tienen sentido se rechazan con un mensaje claro"""
        malos = [
            {"empresa": "E99", "mes": "2036-02", "km_total": 1000},
            {"empresa": "E10", "mes": "no es mes", "km_total": 1000},
            {"empresa": "E10", "mes": "2036-02", "km_total": 0},
            {"empresa": "E10", "mes": "2036-02", "km_total": -5},
            {"empresa": "E10", "mes": "2036-02", "km_total": 1000, "km_padron": -1},
            {"empresa": "E10", "mes": "2036-02", "km_total": 1000, "km_padron": 100, "km_buseton": 100},
        ]
        for cuerpo in malos:
            r = client.put("/api/dpv-km", json=cuerpo, headers=admin_headers)
            assert r.status_code == status.HTTP_400_BAD_REQUEST, cuerpo

    def test_requiere_admin(self, client, tecnico_headers):
        r = client.put("/api/dpv-km", json={"empresa": "E10", "mes": "2036-02", "km_total": 1000}, headers=tecnico_headers)
        assert r.status_code == status.HTTP_403_FORBIDDEN
        assert client.delete("/api/dpv-km/E10/2036-02", headers=tecnico_headers).status_code == status.HTTP_403_FORBIDDEN


class TestNumeros:

    def test_num_o_none_entiende_los_formatos_de_excel_y_de_texto(self):
        assert datos.num_o_none(920000) == 920000.0
        assert datos.num_o_none("920000") == 920000.0
        assert datos.num_o_none("1.234.567,5") == 1234567.5       # es-CO
        assert datos.num_o_none("1,234,567.5") == 1234567.5
        assert datos.num_o_none("12,5") == 12.5
        assert datos.num_o_none("") is None and datos.num_o_none(None) is None
        assert datos.num_o_none("abc") is None and datos.num_o_none(True) is None

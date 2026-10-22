# ═══════════════════════════════════════════════════════════
#  EMANTIX PRO — Rutas de datos operativos
#  DPV, ICO, Disponibilidad de flota, DPV mensual e historial de cargas
#  (endpoints consumidos por CargaDPV.vue, IndicadoresICO.vue y dashboardStore.js)
# ═══════════════════════════════════════════════════════════

import io
import os
import re
import time
import logging
import unicodedata
from datetime import datetime, date, time as dtime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from config import settings          # antes que dpv_reglas: config.py carga el .env (umbrales DPV_*)
import dpv_reglas as reglas
from database import (
    DpvRegistroDB, IcoRegistroDB, DisponibilidadDB, DpvMensualDB, DpvKmMensualDB, CargaHistorialDB,
)
from deps import get_db, get_current_user, get_current_admin_user

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api", tags=["datos"])

MAX_FILE_MB = 10
BATCH = 500
MESES = {
    "ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SEP": 9, "SET": 9, "OCT": 10, "NOV": 11, "DIC": 12,
}
MES_LABEL = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
# Prefijo de móvil → empresa / unidad funcional
PREFIJO_EMPRESA = {"Z32": "E10", "Z34": "E16"}


# ── HELPERS ────────────────────────────────────────────────

def norm_header(h: Any) -> str:
    """'F. Inicio DPV ' → 'F INICIO DPV' (sin tildes, mayúsculas, sin puntuación)."""
    s = unicodedata.normalize("NFKD", str(h or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9]+", " ", s).strip().upper()
    return s


def username(user: Dict) -> str:
    return str(user.get("username") or user.get("sub") or "")


def to_dt(v: Any) -> Optional[datetime]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)):  # serial de Excel
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v)
        except Exception:
            return None
    s = str(v).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d",
                # ISO 8601 con "T" (algunas celdas del Excel llegan como texto, p. ej. 2026-09-21T14:30:00)
                "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S.%f",
                "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M",
                "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def to_mes(v: Any) -> Optional[datetime]:
    """Acepta fecha, '2026-07', '07/2026', 'JULIO 2026', 'jul-26'… → primer día del mes."""
    d = to_dt(v)
    if d:
        return datetime(d.year, d.month, 1)
    s = norm_header(v)
    m = re.match(r"^(\d{4}) (\d{1,2})$", s) or None
    if m:
        return datetime(int(m.group(1)), int(m.group(2)), 1)
    m = re.match(r"^(\d{1,2}) (\d{4})$", s)
    if m:
        return datetime(int(m.group(2)), int(m.group(1)), 1)
    m = re.match(r"^([A-Z]+) (?:DE )?(\d{2,4})$", s)
    if m and m.group(1)[:3] in MESES:
        y = int(m.group(2))
        y = y + 2000 if y < 100 else y
        return datetime(y, MESES[m.group(1)[:3]], 1)
    return None


def to_int(v: Any, default: Optional[int] = 0) -> Optional[int]:
    if v is None or v == "":
        return default
    try:
        return int(float(str(v).replace(",", "").strip()))
    except (ValueError, TypeError):
        return default


def to_float(v: Any) -> float:
    try:
        return float(str(v).replace(",", ".").strip())
    except (ValueError, TypeError):
        return 0.0


def fix_mojibake(s: Any) -> Any:
    """Repara texto con doble codificación UTF-8 (p. ej. 'cÃ³digo' → 'código', 'seÃ±al' → 'señal').

    Cada carácter se convierte de vuelta al byte original (latin-1 / cp1252) y el resultado se
    decodifica como UTF-8. Si el texto no tiene ese problema (o no se puede revertir de forma
    exacta) se devuelve tal cual, por lo que es seguro e idempotente.
    """
    if not isinstance(s, str) or s.isascii():
        return s
    try:
        raw = bytearray()
        for ch in s:
            o = ord(ch)
            if o < 256:
                raw.append(o)
            else:
                raw.extend(ch.encode("cp1252"))
        return raw.decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return s


def to_str(v: Any, maxlen: Optional[int] = None) -> str:
    if v is None:
        return ""
    if isinstance(v, dtime):
        s = v.strftime("%H:%M:%S")
    elif isinstance(v, datetime):
        s = v.isoformat(sep=" ")
    else:
        # Reparar ANTES de strip(): un espacio de no separación (NBSP) final llega como "Â\xa0" y
        # strip() se llevaría el \xa0 dejando el "Â" huérfano, lo que impediría la reparación.
        s = fix_mojibake(str(v)).strip()
    return s[:maxlen] if maxlen else s


def iso(v: Optional[datetime]) -> Optional[str]:
    return v.isoformat() if v else None


def row_to_dict(obj) -> Dict[str, Any]:
    out = {}
    for c in obj.__table__.columns:
        v = getattr(obj, c.name)
        out[c.name] = v.isoformat() if isinstance(v, (datetime, date)) else v
    return out


async def leer_excel(file: UploadFile, solo_xlsx: bool = False):
    nombre = file.filename or ""
    patron = r"\.xlsx$" if solo_xlsx else r"\.(xlsx|xlsm|xls)$"
    if not re.search(patron, nombre, re.I):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "Solo se permiten archivos .xlsx" if solo_xlsx else "Solo se permiten archivos .xlsx o .xls")
    contenido = await file.read()
    if len(contenido) > MAX_FILE_MB * 1024 * 1024:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, f"El archivo supera {MAX_FILE_MB} MB")
    if nombre.lower().endswith(".xls"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "El formato .xls antiguo no es compatible; guarda el archivo como .xlsx")
    from openpyxl import load_workbook
    try:
        return load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"No se pudo leer el Excel: {e}")


def filas_hoja(ws, max_header_scan: int = 10):
    """Devuelve (headers_normalizados, iterador de dicts). Busca la fila de encabezados en las primeras filas."""
    filas = ws.iter_rows(values_only=True)
    headers: List[str] = []
    for _ in range(max_header_scan):
        try:
            r = next(filas)
        except StopIteration:
            return [], iter(())
        cand = [norm_header(c) for c in r]
        if sum(1 for c in cand if c) >= 3:
            headers = cand
            break

    def gen():
        for r in filas:
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            yield {headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]}
    return headers, gen()


def pick(row: Dict[str, Any], *keys: str) -> Any:
    for k in keys:
        if k in row and row[k] not in (None, ""):
            return row[k]
    return None


def faltantes(headers: List[str], requeridas: List[str]) -> List[str]:
    return [r for r in requeridas if r not in headers]


def hoja_con_columnas(wb, requeridas: List[str], preferidas: tuple = ()):
    """Busca la hoja del libro que trae las columnas requeridas (el Excel del DPV tiene 17 hojas y la activa
    puede ser cualquiera). Prueba primero las hojas de nombre «preferido». Devuelve (nombre, headers, filas);
    si ninguna sirve, devuelve la hoja activa para que el error diga qué columnas faltan."""
    hojas = sorted(wb.worksheets, key=lambda w: 0 if norm_header(w.title) in preferidas else 1)
    for ws in hojas:
        headers, filas = filas_hoja(ws)
        if headers and not faltantes(headers, requeridas):
            return ws.title, headers, filas
    ws = wb.active
    headers, filas = filas_hoja(ws)
    return ws.title, headers, filas


def registrar_historial(db: Session, tipo: str, user: Dict, archivo: str, total: int,
                        ins: int, upd: int, err: int, seg: float, msg: str = "") -> None:
    estado = "exitoso" if err == 0 else ("parcial" if ins + upd > 0 else "fallido")
    db.add(CargaHistorialDB(
        tipo_carga=tipo, usuario_id=username(user), nombre_archivo=(archivo or "")[:300],
        total_filas=total, filas_insertadas=ins, filas_actualizadas=upd, filas_error=err,
        estado=estado, mensaje_error=msg[:4000], tiempo_procesamiento_seg=str(round(seg, 2)),
    ))
    db.commit()


def respuesta(total, ins, upd, errores, seg, extra=None):
    r = {
        "message": "Archivo procesado exitosamente" if not errores else
                   (f"Procesamiento parcial: {len(errores)} filas con error" if ins + upd else "No se pudo procesar ninguna fila"),
        "total": total, "insertados": ins, "actualizados": upd,
        "errores": len(errores), "errores_detalle": errores[:20], "tiempo_seg": round(seg, 2),
    }
    if extra:
        r.update(extra)
    return r


def upsert(db: Session, model, key: str, registros: List[Dict[str, Any]]):
    """Inserta o actualiza por la columna `key`, en lotes. Devuelve (insertados, actualizados)."""
    ins = upd = 0
    for i in range(0, len(registros), BATCH):
        lote = {r[key]: r for r in registros[i:i + BATCH]}  # dedup dentro del lote
        col = getattr(model, key)
        existentes = {getattr(o, key): o for o in db.query(model).filter(col.in_(list(lote.keys()))).all()}
        for k, data in lote.items():
            obj = existentes.get(k)
            if obj:
                for campo, v in data.items():
                    setattr(obj, campo, v)
                obj.updated_at = datetime.utcnow()
                upd += 1
            else:
                db.add(model(**data))
                ins += 1
        db.commit()
    return ins, upd


# ═══════════════════════════════════════════════════════════
#  DPV
# ═══════════════════════════════════════════════════════════

DPV_COLS = {
    "dpv_id": ("DPV",), "estado": ("ESTADO",), "fecha_inicio": ("F INICIO DPV", "FECHA INICIO DPV"),
    "fecha_cierre": ("F CIERRE DPV", "FECHA CIERRE DPV"), "empresa": ("EMPRESA",), "fuente": ("FUENTE",),
    "placa": ("PLACA",), "movil": ("MOVIL",), "tipologia": ("TIPOLOGIA",),
    "fecha_inmovilizacion": ("FECHA INMOVILIZACION", "F INMOVILIZACION"), "hora": ("HORA",),
    "causa_inmovilizacion": ("CAUSA DE INMOVILIZACION", "CAUSA INMOVILIZACION"),
    "descripcion_novedad": ("DESCRIPCION DE LA NOVEDAD", "DESCRIPCION NOVEDAD", "DESCRIPCION"),
    "ruta": ("RUTA",), "operador": ("OPERADOR",), "dia_semana": ("DIA SEMANA", "DIA DE LA SEMANA"),
    "alerta": ("ALERTA",), "franja_horaria": ("FRANJA HORARIA", "FRANJA"),
}
DPV_LEN = {"estado": 100, "empresa": 200, "fuente": 200, "placa": 20, "movil": 20, "tipologia": 100,
           "hora": 10, "causa_inmovilizacion": 300, "ruta": 100, "operador": 50, "dia_semana": 20,
           "alerta": 100, "franja_horaria": 50}


DPV_HOJAS = ("BASE DE DATOS", "BASE", "DPV", "EVENTOS")   # nombres de hoja que se prueban primero


@router.post("/dpv/upload")
async def upload_dpv(file: UploadFile = File(...), db: Session = Depends(get_db),
                     user: Dict = Depends(get_current_admin_user)):
    t0 = time.time()
    wb = await leer_excel(file)
    # El libro del DPV trae varias hojas (base de datos, km por día, DPV_E10…): se usa la que tiene las columnas del evento
    hoja, headers, filas = hoja_con_columnas(wb, ["DPV", "ESTADO", "PLACA"], DPV_HOJAS)
    falta = faltantes(headers, ["DPV", "ESTADO", "PLACA"])
    if falta:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            f"Faltan columnas requeridas: {', '.join(falta)} (hoja revisada: «{hoja}»)")

    registros, errores, total = [], [], 0
    calidad: Dict[str, int] = {}      # cuántas filas tuvieron cada arreglo (espacios, estado, tipología, hora…)
    reparados = 0                     # textos con doble codificación (Ã³ → ó) que se repararon
    for n, row in enumerate(filas, start=2):
        total += 1
        dpv_id = to_int(pick(row, "DPV"), None)
        if dpv_id is None:
            errores.append(f"Fila {n}: DPV vacío o no numérico")
            continue
        r: Dict[str, Any] = {"dpv_id": dpv_id}
        for campo, keys in DPV_COLS.items():
            if campo == "dpv_id":
                continue
            v = pick(row, *keys)
            if campo.startswith("fecha"):
                r[campo] = to_dt(v)
                continue
            if isinstance(v, float) and v.is_integer():
                v = int(v)                                   # operador 320000.0 → 320000
            if isinstance(v, str) and fix_mojibake(v) != v:
                reparados += 1
            if campo == "hora":
                r[campo] = to_str(v.time() if isinstance(v, datetime) else v, 10)
            else:
                r[campo] = to_str(v, DPV_LEN.get(campo))
        r, cambios = reglas.limpiar_evento(r)
        for campo in cambios:
            calidad[campo] = calidad.get(campo, 0) + 1
        registros.append(r)

    try:
        ins, upd = upsert(db, DpvRegistroDB, "dpv_id", registros)
    except Exception as e:
        db.rollback()
        logger.error("Error guardando DPV", exc_info=True)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"Error guardando en BD: {e}")

    seg = time.time() - t0
    registrar_historial(db, "dpv", user, file.filename, total, ins, upd, len(errores), seg, "\n".join(errores[:50]))
    preview = [{"dpv": r["dpv_id"], "estado": r["estado"], "placa": r["placa"], "empresa": r["empresa"],
                "causa": r["causa_inmovilizacion"],
                "fecha": r["fecha_inmovilizacion"].strftime("%Y-%m-%d") if r["fecha_inmovilizacion"] else ""}
               for r in registros[:10]]
    descartadas = sum(1 for r in registros if not reglas.es_inmovilizacion(r["estado"]))
    return respuesta(total, ins, upd, errores, seg, {
        "preview": preview, "hoja": hoja,
        "calidad": {"textos_reparados": reparados, "normalizados": calidad,
                    "no_cuentan_para_dpv": descartadas},
    })

# Columnas que realmente usa el dashboard (evita enviar created_at, fuente, etc.)
DPV_SLIM = ("id", "dpv_id", "estado", "fecha_inicio", "fecha_cierre", "empresa", "placa", "movil", "tipologia",
            "fecha_inmovilizacion", "hora", "causa_inmovilizacion", "descripcion_novedad", "ruta",
            "operador", "dia_semana", "alerta", "franja_horaria")


@router.get("/dpv")
def listar_dpv(limit: int = Query(100, ge=1, le=100000), offset: int = Query(0, ge=0),
               mes: Optional[str] = None, slim: bool = False,
               db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    q = db.query(DpvRegistroDB)

    # Filtro por mes ('YYYY-MM'): usa to_mes(), que ya entiende ese formato
    if mes:
        d = to_mes(mes)
        if d:
            fin = datetime(d.year + (1 if d.month == 12 else 0), d.month % 12 + 1, 1)
            q = q.filter(DpvRegistroDB.fecha_inmovilizacion >= d,
                         DpvRegistroDB.fecha_inmovilizacion < fin)

    total = q.count()
    orden = (DpvRegistroDB.fecha_inmovilizacion.desc(), DpvRegistroDB.id.desc())

    if slim:
        cols = [getattr(DpvRegistroDB, c) for c in DPV_SLIM]
        filas = q.with_entities(*cols).order_by(*orden).offset(offset).limit(limit).all()
        items = [
            {c: (v.isoformat() if isinstance(v, (datetime, date)) else v)
             for c, v in zip(DPV_SLIM, f)}
            for f in filas
        ]
    else:
        items = [row_to_dict(i) for i in q.order_by(*orden).offset(offset).limit(limit).all()]

    # La regla «¿cuenta como inmovilización?» vive en el backend (dpv_reglas): el tablero la lee, no la reimplementa
    for it in items:
        it["cuenta"] = reglas.es_inmovilizacion(it.get("estado"))

    return {"total": total, "items": items}

# ═══════════════════════════════════════════════════════════
#  ICO
# ═══════════════════════════════════════════════════════════

ICO_COLS = {
    "id_novedad": ("ID NOVEDAD",), "estado": ("ESTADO",),
    "fecha_inicio": ("F INICIO DPV", "F INICIO ICO", "F INICIO", "FECHA INICIO"),
    "fecha_cierre": ("F CIERRE DPV", "F CIERRE ICO", "F CIERRE", "FECHA CIERRE"),
    "empresa": ("EMPRESA",), "tipo_novedad": ("TIPO NOVEDAD", "TIPO DE NOVEDAD"),
    "fecha_novedad": ("F NOVEDAD", "FECHA NOVEDAD"), "fecha_identificacion": ("F IDENTIFICACION", "FECHA IDENTIFICACION"),
    "fecha_notificacion": ("F NOTIFICACION", "FECHA NOTIFICACION"), "fuente": ("FUENTE",), "area": ("AREA",),
    "direccion": ("DIRECCION",), "placa": ("PLACA",), "movil": ("MOVIL",), "tipologia": ("TIPOLOGIA",),
    "dia_semana": ("DIA SEMANA", "DIA DE LA SEMANA"), "puntos": ("PUNTOS",), "descripcion": ("DESCRIPCION",),
}
ICO_LEN = {"estado": 100, "empresa": 200, "tipo_novedad": 100, "fuente": 200, "area": 100, "direccion": 300,
           "placa": 20, "movil": 20, "tipologia": 100, "dia_semana": 20, "puntos": 50}


@router.post("/ico/upload")
async def upload_ico(file: UploadFile = File(...), db: Session = Depends(get_db),
                     user: Dict = Depends(get_current_admin_user)):
    t0 = time.time()
    wb = await leer_excel(file)
    headers, filas = filas_hoja(wb.active)
    falta = faltantes(headers, ["ICO", "ESTADO", "PLACA"])
    if falta:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Faltan columnas requeridas: {', '.join(falta)}")

    registros, errores, total = [], [], 0
    for n, row in enumerate(filas, start=2):
        total += 1
        ico_id = to_int(pick(row, "ICO"), None)
        if ico_id is None:
            errores.append(f"Fila {n}: ICO vacío o no numérico")
            continue
        r: Dict[str, Any] = {"ico_id": ico_id}
        for campo, keys in ICO_COLS.items():
            v = pick(row, *keys)
            if campo.startswith("fecha"):
                r[campo] = to_dt(v)
            elif campo == "id_novedad":
                r[campo] = to_int(v, None)
            elif campo == "puntos":
                # Entero normalizado ("10", no "10.0"). Vacío = sin puntos asignados (no es 0).
                p = to_int(v, None)
                r[campo] = "" if p is None else str(p)
            else:
                r[campo] = to_str(v, ICO_LEN.get(campo))
        registros.append(r)

    try:
        ins, upd = upsert(db, IcoRegistroDB, "ico_id", registros)
    except Exception as e:
        db.rollback()
        logger.error("Error guardando ICO", exc_info=True)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"Error guardando en BD: {e}")

    seg = time.time() - t0
    registrar_historial(db, "ico", user, file.filename, total, ins, upd, len(errores), seg, "\n".join(errores[:50]))
    return respuesta(total, ins, upd, errores, seg)


def tipo_infraccion(tipo_novedad: Optional[str]) -> str:
    """Clasificación de la infracción (columna calculada del Power BI).

    "Directo" para los códigos configurados en ICO_CODIGOS_DIRECTO (por defecto I5006, I5013,
    I5022 e I5023); cualquier otro código es "Hallazgo Rechazado".
    """
    codigo = (tipo_novedad or "").strip().upper()
    return "Directo" if codigo in settings.ico_codigos_directo else "Hallazgo Rechazado"


def _filtrar_ico(q, fecha_desde, fecha_hasta, empresa, tipo_novedad, estado, placa):
    if fecha_desde and to_dt(fecha_desde):
        q = q.filter(IcoRegistroDB.fecha_novedad >= to_dt(fecha_desde))
    if fecha_hasta and to_dt(fecha_hasta):
        hasta = to_dt(fecha_hasta).replace(hour=23, minute=59, second=59)
        q = q.filter(IcoRegistroDB.fecha_novedad <= hasta)
    if empresa:
        q = q.filter(IcoRegistroDB.empresa == empresa)
    if tipo_novedad:
        q = q.filter(IcoRegistroDB.tipo_novedad == tipo_novedad)
    if estado:
        q = q.filter(IcoRegistroDB.estado == estado)
    if placa:
        q = q.filter(IcoRegistroDB.placa.ilike(f"%{placa.strip()}%"))
    return q


@router.get("/ico/resumen")
def resumen_ico(db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    total = db.query(func.count(IcoRegistroDB.id)).scalar() or 0
    cnt = func.count(IcoRegistroDB.id)
    por_tipo = db.query(IcoRegistroDB.tipo_novedad, cnt).group_by(IcoRegistroDB.tipo_novedad).order_by(cnt.desc()).all()
    por_estado = db.query(IcoRegistroDB.estado, cnt).group_by(IcoRegistroDB.estado).order_by(cnt.desc()).all()
    top_placas = db.query(IcoRegistroDB.placa, cnt).group_by(IcoRegistroDB.placa).order_by(cnt.desc()).limit(10).all()
    return {
        "total": total,
        "por_tipo": [{"tipo": t, "total": n} for t, n in por_tipo],
        "por_estado": [{"estado": e, "total": n} for e, n in por_estado],
        "top_placas": [{"placa": p, "total": n} for p, n in top_placas],
    }


@router.get("/ico/filtros")
def filtros_ico(db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    def distintos(col):
        return sorted(v for (v,) in db.query(col).distinct().all() if v)
    return {
        "empresas": distintos(IcoRegistroDB.empresa),
        "tipos": distintos(IcoRegistroDB.tipo_novedad),
        "estados": distintos(IcoRegistroDB.estado),
    }


@router.get("/ico")
def listar_ico(limit: int = Query(50, ge=1, le=100000), offset: int = Query(0, ge=0),
               fecha_desde: Optional[str] = None, fecha_hasta: Optional[str] = None,
               empresa: Optional[str] = None, tipo_novedad: Optional[str] = None,
               estado: Optional[str] = None, placa: Optional[str] = None,
               db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    q = _filtrar_ico(db.query(IcoRegistroDB), fecha_desde, fecha_hasta, empresa, tipo_novedad, estado, placa)
    total = q.count()
    items = q.order_by(IcoRegistroDB.fecha_novedad.desc(), IcoRegistroDB.id.desc()).offset(offset).limit(limit).all()
    out = []
    for i in items:
        d = row_to_dict(i)
        d["fecha_novedad"] = i.fecha_novedad.strftime("%Y-%m-%d") if i.fecha_novedad else ""
        d["hora"] = i.fecha_novedad.strftime("%H:%M:%S") if i.fecha_novedad else ""
        # Puntos como número (los registros viejos se guardaron como "10.0"); None = sin puntos.
        d["puntos"] = to_int(i.puntos, None)
        d["tipo_infraccion"] = tipo_infraccion(i.tipo_novedad)
        out.append(d)
    return {"total": total, "items": out}


# ═══════════════════════════════════════════════════════════
#  DISPONIBILIDAD DE FLOTA
# ═══════════════════════════════════════════════════════════

def empresa_por_movil(movil: str) -> Optional[str]:
    m = (movil or "").strip().upper()
    for pref, emp in PREFIJO_EMPRESA.items():
        if m.startswith(pref):
            return emp
    return None


@router.post("/disponibilidad/upload")
async def upload_disponibilidad(file: UploadFile = File(...), db: Session = Depends(get_db),
                                user: Dict = Depends(get_current_admin_user)):
    """Cada archivo es una foto (corte) de una unidad: reemplaza el corte anterior de esa unidad."""
    t0 = time.time()
    wb = await leer_excel(file)
    headers, filas = filas_hoja(wb.active)
    falta = faltantes(headers, ["MOVIL"])
    if falta:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Falta la columna requerida: MOVIL")

    registros, errores, total = [], [], 0
    for n, row in enumerate(filas, start=2):
        total += 1
        movil = to_str(pick(row, "MOVIL"), 50)
        emp = empresa_por_movil(movil)
        if not emp:
            errores.append(f"Fila {n}: móvil '{movil}' no empieza por Z32- ni Z34-")
            continue
        inm = pick(row, "INMOVILIZADO", "ESTADO")
        registros.append({
            "empresa": emp,
            "movil": movil,
            "fecha_corte": to_dt(pick(row, "FECHA DE CORTE", "FECHA CORTE")),
            "fecha_ingreso": to_dt(pick(row, "FECHA DE INGRESO A MANTENIMIENTO", "FECHA INGRESO A MANTENIMIENTO",
                                        "FECHA DE INGRESO", "FECHA INGRESO")),
            "dias_inoperatividad": to_float(pick(row, "DIAS DE INOPERATIVIDAD", "DIAS INOPERATIVIDAD", "DIAS")),
            "area": to_str(pick(row, "AREA", "AREA RESPONSABLE", "SISTEMA"), 150) or "Sin área",
            "descripcion": to_str(pick(row, "DESCRIPCION DE LA FALLA", "DESCRIPCION FALLA", "DESCRIPCION")),
            "inmovilizado": norm_header(inm) in ("SI", "S", "X", "TRUE", "1", "INMOVILIZADO"),
        })

    if not registros:
        seg = time.time() - t0
        registrar_historial(db, "disponibilidad", user, file.filename, total, 0, 0, len(errores), seg)
        return respuesta(total, 0, 0, errores or ["El archivo no tiene filas válidas"], seg)

    empresas = sorted({r["empresa"] for r in registros})
    try:
        reemplazados = db.query(DisponibilidadDB).filter(DisponibilidadDB.empresa.in_(empresas)) \
                         .delete(synchronize_session=False)
        db.bulk_insert_mappings(DisponibilidadDB, registros)
        db.commit()
    except Exception as e:
        db.rollback()
        logger.error("Error guardando disponibilidad", exc_info=True)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"Error guardando en BD: {e}")

    seg = time.time() - t0
    registrar_historial(db, "disponibilidad", user, file.filename, total, len(registros), 0, len(errores), seg)
    r = respuesta(total, len(registros), 0, errores, seg, {"unidades": empresas, "reemplazados": reemplazados})
    r["message"] = f"Disponibilidad {', '.join(empresas)} actualizada" + (f" ({len(errores)} filas con error)" if errores else "")
    return r


def _flota_total(db: Session, emp: str) -> Optional[int]:
    env = os.getenv(f"FLOTA_TOTAL_{emp}")
    if env and env.strip().isdigit():
        return int(env)
    # Respaldo: móviles distintos vistos en DPV con el prefijo de la unidad
    pref = next(p for p, e in PREFIJO_EMPRESA.items() if e == emp)
    n = db.query(func.count(func.distinct(DpvRegistroDB.movil))) \
          .filter(DpvRegistroDB.movil.like(f"{pref}%")).scalar()
    return n or None


@router.get("/disponibilidad")
def obtener_disponibilidad(db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    resultado: Dict[str, Any] = {}
    for emp in ("E10", "E16"):
        filas = db.query(DisponibilidadDB).filter(DisponibilidadDB.empresa == emp).all()
        if not filas:
            resultado[emp.lower()] = None
            continue
        cortes = [f.fecha_corte for f in filas if f.fecha_corte]
        no_disp = len({f.movil for f in filas})
        flota = _flota_total(db, emp)
        if flota is not None:
            flota = max(flota, no_disp)
        disponibles = (flota - no_disp) if flota else None
        por_area: Dict[str, int] = {}
        for f in filas:
            por_area[f.area or "Sin área"] = por_area.get(f.area or "Sin área", 0) + 1
        resultado[emp.lower()] = {
            "fecha_corte": max(cortes).strftime("%Y-%m-%d") if cortes else None,
            "flota_total": flota,
            "no_disponibles": no_disp,
            "disponibles": disponibles,
            "porcentaje": round(disponibles / flota * 100, 1) if flota else None,
            "por_area": [{"area": a, "total": n} for a, n in sorted(por_area.items(), key=lambda x: -x[1])],
            "detalle": [{
                "movil": f.movil, "area": f.area, "descripcion": f.descripcion,
                "dias_inoperatividad": round(f.dias_inoperatividad or 0, 1),
                "inmovilizado": bool(f.inmovilizado),
                "fecha_ingreso": f.fecha_ingreso.strftime("%Y-%m-%d") if f.fecha_ingreso else None,
            } for f in filas],
        }
    return resultado


# ═══════════════════════════════════════════════════════════
#  DPV MENSUAL (kilometraje)
#  Los km oficiales de cada mes van en dpv_km_mensual (carga del Excel o captura manual). Las inmovilizaciones y el
#  DPV se calculan SIEMPRE en vivo desde los eventos (dpv_registros) con dpv_reglas.py: si un evento pasa a
#  «Contestado Contundente», el DPV del mes se actualiza solo.
# ═══════════════════════════════════════════════════════════

def num_o_none(v: Any) -> Optional[float]:
    """Número de una celda: int/float tal cual; texto '1.234.567,5' (es-CO) o '1,234,567.5'. Vacío o inválido → None."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(" ", "")
    if not t:
        return None
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", t):
        t = t.replace(",", "")
    else:
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


@router.post("/dpv-mensual/upload")
async def upload_dpv_mensual(file: UploadFile = File(...), db: Session = Depends(get_db),
                             user: Dict = Depends(get_current_admin_user)):
    """Hojas DPV_E10 y DPV_E16. Entiende el diseño nuevo (MES, TOTAL KM, KM PADRON, KM BUSETON, TOTAL INNMOVILIZADOS…)
    y el anterior (DVP_E10, No VARADOS, KMS…). Los meses sin km (p. ej. los que aún no se han cerrado) se omiten:
    no son cero kilómetros."""
    t0 = time.time()
    wb = await leer_excel(file, solo_xlsx=True)
    hojas = {norm_header(n).replace(" ", "_"): n for n in wb.sheetnames}
    faltan = [h for h in ("DPV_E10", "DPV_E16") if h not in hojas and h.replace("DPV", "DVP") not in hojas]
    if faltan:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            f"Faltan las hojas: {', '.join(faltan)} (encontradas: {', '.join(wb.sheetnames)})")

    ins = upd = total = 0
    errores: List[str] = []
    advertencias: List[str] = []
    sin_km: Dict[str, List[str]] = {"E10": [], "E16": []}
    formatos = set()
    for emp in ("E10", "E16"):
        nombre = hojas.get(f"DPV_{emp}") or hojas.get(f"DVP_{emp}")
        headers, filas = filas_hoja(wb[nombre])
        nuevo = "TOTAL KM" in headers or "KM TOTAL" in headers
        formatos.add("nuevo" if nuevo else "anterior")
        existentes = ({r.mes: r for r in db.query(DpvKmMensualDB).filter(DpvKmMensualDB.empresa == emp).all()}
                      if nuevo else {})
        for n, row in enumerate(filas, start=2):
            total += 1
            mes = to_mes(pick(row, "MES", "FECHA"))
            if not mes:
                errores.append(f"Hoja {nombre}, fila {n}: MES/FECHA no reconocido")
                continue
            etiqueta = mes.strftime("%Y-%m")

            if nuevo:
                km = num_o_none(pick(row, "TOTAL KM", "KM TOTAL"))
                if not km or km <= 0:
                    sin_km[emp].append(etiqueta)
                    continue
                km_p, km_b = num_o_none(pick(row, "KM PADRON")), num_o_none(pick(row, "KM BUSETON"))
                if km_p is not None and km_b is not None and abs(km_p + km_b - km) > max(1.0, km * 0.01):
                    advertencias.append(f"{emp} {etiqueta}: KM PADRON + KM BUSETON ({km_p + km_b:,.0f}) no coincide con TOTAL KM ({km:,.0f})")
                datos = {
                    "km_total": km, "km_padron": km_p, "km_buseton": km_b,
                    "inm_excel": to_int(pick(row, "TOTAL INNMOVILIZADOS", "TOTAL INMOVILIZADOS", "TOTAL INNMOVILIZADO",
                                             "TOTAL INMOVILIZADO", "NO VARADOS"), None),
                    "dpv_excel": num_o_none(pick(row, "DPV GENERAL", f"DPV {emp}", f"DVP {emp}")),
                    "fuente": "excel", "updated_by": username(user),
                }
                obj = existentes.get(mes)
                if obj:
                    for k, v in datos.items():
                        setattr(obj, k, v)
                    obj.updated_at = datetime.utcnow()
                    upd += 1
                else:
                    obj = DpvKmMensualDB(empresa=emp, mes=mes, **datos)
                    db.add(obj)
                    existentes[mes] = obj
                    ins += 1
                continue

            # ── Diseño anterior: se guarda tal cual en dpv_mensual (solo si el mes trae algún dato) ──
            datos = {
                "fecha_texto": to_str(pick(row, "FECHA"), 100),
                "dpv_general": to_int(pick(row, f"DVP {emp}", f"DPV {emp}", "DPV GENERAL", "DPV")),
                "dpv_estandar": to_int(pick(row, "DPV ESTANDAR"), 10000),
                "dpv_critico": to_int(pick(row, "DPV CRITICO"), 7000),
                "no_varados": to_int(pick(row, "NO VARADOS", "N VARADOS", "VARADOS")),
                "kms": to_int(pick(row, "KMS", "KM")),
                "padron": to_int(pick(row, "PADRON")),
                "buseton": to_int(pick(row, "BUSETON")),
            }
            if not datos["dpv_general"] and not datos["kms"]:
                sin_km[emp].append(etiqueta)
                continue
            obj = db.query(DpvMensualDB).filter(DpvMensualDB.empresa == emp, DpvMensualDB.mes == mes).first()
            if obj:
                for k, v in datos.items():
                    setattr(obj, k, v)
                obj.updated_at = datetime.utcnow()
                upd += 1
            else:
                db.add(DpvMensualDB(empresa=emp, mes=mes, **datos))
                ins += 1
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        logger.error("Error guardando DPV mensual", exc_info=True)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"Error guardando en BD: {e}")

    seg = time.time() - t0
    registrar_historial(db, "dpv_mensual", user, file.filename, total, ins, upd, len(errores), seg, "\n".join(errores[:50]))
    return respuesta(total, ins, upd, errores, seg, {
        "formato": "nuevo" if formatos == {"nuevo"} else "anterior" if formatos == {"anterior"} else "mixto",
        "sin_km": sin_km, "advertencias": advertencias[:20],
    })


def _serie_dpv(db: Session) -> Dict[str, Any]:
    """Serie mensual por unidad: km cargados + inmovilizaciones calculadas en vivo desde los eventos."""
    filas = db.query(DpvRegistroDB.fecha_inmovilizacion, DpvRegistroDB.empresa, DpvRegistroDB.movil,
                     DpvRegistroDB.tipologia, DpvRegistroDB.estado) \
              .filter(DpvRegistroDB.fecha_inmovilizacion.isnot(None)).all()
    mapa = reglas.agregar_inmovilizaciones(filas)
    km: Dict[str, List[Dict[str, Any]]] = {"E10": [], "E16": []}
    for r in db.query(DpvKmMensualDB).order_by(DpvKmMensualDB.mes.asc()).all():
        if r.empresa in km:
            km[r.empresa].append({
                "mes": r.mes.strftime("%Y-%m"), "km_total": r.km_total, "km_padron": r.km_padron,
                "km_buseton": r.km_buseton, "inm_excel": r.inm_excel, "dpv_excel": r.dpv_excel,
                "fuente": r.fuente or "excel",
            })
    legado: Dict[str, List[Dict[str, Any]]] = {"E10": [], "E16": []}
    for r in db.query(DpvMensualDB).order_by(DpvMensualDB.mes.asc()).all():
        if r.empresa in legado:
            legado[r.empresa].append({"mes": r.mes.strftime("%Y-%m"), "dpv": r.dpv_general,
                                      "kms": r.kms, "varados": r.no_varados})
    return {
        "e10": reglas.serie_mensual("E10", km["E10"], mapa, legado["E10"]),
        "e16": reglas.serie_mensual("E16", km["E16"], mapa, legado["E16"]),
        "umbrales": {"estandar": reglas.ESTANDAR_KM, "critico": reglas.CRITICO_KM},
        "estados_descartados": list(reglas.ESTADOS_DESCARTADOS_TXT),     # estados que NO cuentan como inmovilización
    }


@router.get("/dpv-mensual")
def obtener_dpv_mensual(db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    return _serie_dpv(db)


class KmMensualIn(BaseModel):
    empresa: str                          # 'E10' | 'E16'
    mes: str                              # 'YYYY-MM'
    km_total: float
    km_padron: Optional[float] = None
    km_buseton: Optional[float] = None


@router.put("/dpv-km")
def guardar_km_mensual(body: KmMensualIn, db: Session = Depends(get_db),
                       user: Dict = Depends(get_current_admin_user)):
    """Captura manual del km oficial de un mes (útil para los meses que el Excel aún no trae)."""
    t0 = time.time()
    emp = (body.empresa or "").strip().upper()
    if emp not in reglas.UNIDADES:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "La unidad debe ser E10 o E16")
    mes = to_mes(body.mes)
    if not mes:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Mes no válido (usa AAAA-MM)")
    total = body.km_total
    if not total or total <= 0:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "El km total debe ser mayor que 0")
    for v in (body.km_padron, body.km_buseton):
        if v is not None and v < 0:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Los km por tipología no pueden ser negativos")
    if body.km_padron is not None and body.km_buseton is not None \
            and abs(body.km_padron + body.km_buseton - total) > max(1.0, total * 0.01):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            f"Padrón + Buseton ({body.km_padron + body.km_buseton:,.0f}) no coincide con el total ({total:,.0f})")

    obj = db.query(DpvKmMensualDB).filter(DpvKmMensualDB.empresa == emp, DpvKmMensualDB.mes == mes).first()
    campos = {"km_total": total, "km_padron": body.km_padron, "km_buseton": body.km_buseton,
              "fuente": "manual", "updated_by": username(user)}
    if obj:
        for k, v in campos.items():
            setattr(obj, k, v)
        obj.updated_at = datetime.utcnow()
    else:
        db.add(DpvKmMensualDB(empresa=emp, mes=mes, **campos))
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        logger.error("Error guardando km mensual", exc_info=True)
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"Error guardando en BD: {e}")
    registrar_historial(db, "dpv_mensual", user, f"Captura manual {emp} {mes.strftime('%Y-%m')}", 1,
                        0 if obj else 1, 1 if obj else 0, 0, time.time() - t0)
    return {"empresa": emp, "mes": mes.strftime("%Y-%m"), "actualizado": bool(obj)}


@router.delete("/dpv-km/{empresa}/{mes}")
def borrar_km_mensual(empresa: str, mes: str, db: Session = Depends(get_db),
                      user: Dict = Depends(get_current_admin_user)):
    emp, m = (empresa or "").strip().upper(), to_mes(mes)
    if emp not in reglas.UNIDADES or not m:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Unidad o mes no válidos")
    n = db.query(DpvKmMensualDB).filter(DpvKmMensualDB.empresa == emp, DpvKmMensualDB.mes == m) \
          .delete(synchronize_session=False)
    db.commit()
    return {"eliminado": bool(n)}


# ═══════════════════════════════════════════════════════════
#  HISTORIAL DE CARGAS
# ═══════════════════════════════════════════════════════════

@router.get("/cargas/historial")
def historial_cargas(tipo: Optional[str] = None, limit: int = Query(10, ge=1, le=200),
                     db: Session = Depends(get_db), user: Dict = Depends(get_current_user)):
    q = db.query(CargaHistorialDB)
    if tipo:
        q = q.filter(CargaHistorialDB.tipo_carga == tipo)
    return [{
        "id": h.id,
        "fecha": h.created_at.strftime("%Y-%m-%d %H:%M") if h.created_at else "",
        "usuario": h.usuario_id, "archivo": h.nombre_archivo, "total": h.total_filas,
        "insertados": h.filas_insertadas, "actualizados": h.filas_actualizadas, "errores": h.filas_error,
        "tiempo_seg": h.tiempo_procesamiento_seg, "estado": h.estado,
    } for h in q.order_by(CargaHistorialDB.created_at.desc()).limit(limit).all()]

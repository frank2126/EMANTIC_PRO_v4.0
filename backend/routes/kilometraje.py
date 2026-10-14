# Kilometraje y programación de móviles  ·  prefijo /api/kilometraje
# Carga la exportación de telemetría (Recorrido / Mín / Máx por día), calcula km por odómetro
# y gestiona hitos de mantenimiento: pendiente → programado → intervenido.
import csv, io, json, logging, re, time
from datetime import date, timedelta
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel
from sqlalchemy import func
from sqlalchemy.orm import Session

from database import AppConfigDB, CargaHistorialDB, LecturaKmDB, MovilBaseDB, ProgramacionMovilDB
from deps import get_current_admin_user, get_current_user, get_db

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/kilometraje", tags=["kilometraje"])

# Configuración. Se puede sobreescribir guardando JSON en app_config con key 'km_config'.
CFG = {"prefijos": {"Z3240": "VW", "Z3241": "VW", "Z3440": "VW", "Z3270": "Scania", "Z3271": "Scania", "Z3470": "Scania", "Z3471": "Scania"},
       "paso": {"VW": 25000, "Scania": 10000},   # km entre hitos por grupo
       "salto_km": 100, "bajo_uso": 50, "dias_sin_dato": 3}
HDR = re.compile(r"^(\d{4}-\d{2}-\d{2})\s*\((Recorrido|Mín|Máx)\)$")

def _cfg(db):
    c, r = dict(CFG), db.query(AppConfigDB).filter(AppConfigDB.key == "km_config").first()
    if r and r.value:
        try: c.update(json.loads(r.value))
        except ValueError: logger.warning("km_config inválido; uso valores por defecto")
    return c

def _who(u):
    """Usuario que actúa: tu get_current_admin_user devuelve un dict; también acepta objetos."""
    g = (lambda k: u.get(k)) if isinstance(u, dict) else (lambda k: getattr(u, k, None))
    return str(g("username") or g("sub") or g("name") or "")
def _norm(v): return re.sub(r"[^A-Z0-9]", "", str(v or "").upper())
def _grupo(mv, c): return c["prefijos"].get(mv[:5], "Sin clasificar")
def _num(v):
    try: return float(str(v).replace(",", ".")) if v not in (None, "") else 0.0
    except ValueError: return 0.0
def _fecha(v):
    if hasattr(v, "strftime"): return v.strftime("%Y-%m-%d")
    s = str(v or "").strip(); m = re.match(r"^\d{4}-\d{2}-\d{2}", s)
    if m: return m.group(0)
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})$", s)          # d/m/aaaa (formato Colombia)
    return f"{m.group(3) if len(m.group(3)) == 4 else '20' + m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else ""
def _rows(f: UploadFile):
    raw, name = f.file.read(), (f.filename or "").lower()
    if name.endswith(".csv"):
        t = raw.decode("utf-8-sig", errors="replace")
        return list(csv.reader(io.StringIO(t), delimiter=";" if t.count(";") >= t.count(",") else ","))
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        return [list(r) for r in load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active.iter_rows(values_only=True)]
    raise HTTPException(400, "Solo se aceptan .xlsx o .csv")

# ── Carga de kilometraje ───────────────────────────────────
@router.post("/upload")
def upload_km(files: List[UploadFile] = File(...), db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    t0, datos = time.time(), {}
    for f in files:
        rows = _rows(f)
        ix = {}
        for i, c in enumerate(rows[0] if rows else []):
            m = HDR.match(str(c or "").strip())
            if m: ix.setdefault(m.group(1), {})[m.group(2)] = i
        if not ix: raise HTTPException(400, f"{f.filename}: no encontré columnas 'AAAA-MM-DD (Recorrido|Mín|Máx)'")
        for r in rows[1:]:
            mv = _norm(r[0]) if r else ""
            if not mv: continue
            for d, c in ix.items():
                g = lambda k: _num(r[c[k]]) if k in c and c[k] < len(r) else 0.0
                datos[(mv, d)] = (g("Recorrido"), g("Mín"), g("Máx"))      # el último archivo gana
    fechas = sorted({d for _, d in datos})
    exist = {(l.movil, l.fecha): l for l in db.query(LecturaKmDB).filter(LecturaKmDB.fecha.between(fechas[0], fechas[-1]))}
    ins = upd = 0; nuevos = []
    for (mv, d), (rec, mn, mx) in datos.items():
        l = exist.get((mv, d))
        if l is None: nuevos.append(dict(movil=mv, fecha=d, recorrido=rec, odo_min=mn, odo_max=mx)); ins += 1
        elif abs(l.recorrido - rec) > .005 or abs(l.odo_min - mn) > .005 or abs(l.odo_max - mx) > .005:
            l.recorrido, l.odo_min, l.odo_max = rec, mn, mx; upd += 1
    if nuevos: db.bulk_insert_mappings(LecturaKmDB, nuevos)
    base, ult = {x[0] for x in db.query(MovilBaseDB.movil).all()}, {}
    for (mv, d), (_, _, mx) in sorted(datos.items(), key=lambda kv: kv[0][1]):
        if mx > 0: ult[mv] = mx
    for mv, odo in ult.items():
        if mv not in base: db.add(MovilBaseDB(movil=mv, odo_base=odo))
    db.add(CargaHistorialDB(tipo_carga="km", usuario_id=_who(user), nombre_archivo=", ".join(f.filename or "" for f in files)[:300],
                            total_filas=len(datos), filas_insertadas=ins, filas_actualizadas=upd, filas_error=0, estado="exitoso",
                            tiempo_procesamiento_seg=str(round(time.time() - t0, 2))))
    db.commit()
    return {"moviles": len({m for m, _ in datos}), "dias": len(fechas), "desde": fechas[0], "hasta": fechas[-1], "insertadas": ins, "actualizadas": upd}

# ── Cálculo ────────────────────────────────────────────────
def _calc(ls, desde, hasta, c):
    first = prev = last = None; dd = sal = 0
    for l in ls:
        if l.fecha < desde or l.odo_max <= 0: continue          # Mín = Máx = 0 → sin dato (no es 0 km)
        dd += 1
        if prev and l.odo_min - prev.odo_max > c["salto_km"] + 300 * ((date.fromisoformat(l.fecha) - date.fromisoformat(prev.fecha)).days - 1): sal += 1
        first = first or l; prev = last = l
    nd = (date.fromisoformat(hasta) - date.fromisoformat(desde)).days + 1
    if not first: return {"km": 0.0, "odo": 0.0, "kd": 0.0, "dd": 0, "sd": nd, "sal": 0}
    km = last.odo_max - first.odo_min                              # distancia = odómetro final − inicial (no sumar 'Recorrido')
    return {"km": km, "odo": last.odo_max, "kd": km / ((date.fromisoformat(last.fecha) - date.fromisoformat(first.fecha)).days + 1), "dd": dd, "sd": nd - dd, "sal": sal}

def _pend(db, mv, c):
    paso, b = c["paso"].get(_grupo(mv, c)), db.query(MovilBaseDB).filter(MovilBaseDB.movil == mv).first()
    if not paso or not b: return 0
    hecho = db.query(func.max(ProgramacionMovilDB.hito_km)).filter(ProgramacionMovilDB.movil == mv, ProgramacionMovilDB.estado == "Intervenido").scalar() or 0
    return (int(max(b.odo_base, hecho) // paso) + 1) * paso     # siguiente múltiplo del paso tras lo ya hecho

def _flota(db, c, hz):
    hasta = db.query(func.max(LecturaKmDB.fecha)).scalar()
    if not hasta: return None
    h = date.fromisoformat(hasta); mes0, d14 = h.replace(day=1).isoformat(), (h - timedelta(days=13)).isoformat()
    lect = {}
    for l in db.query(LecturaKmDB).filter(LecturaKmDB.fecha >= min(mes0, d14)).order_by(LecturaKmDB.fecha):
        lect.setdefault(l.movil, []).append(l)
    base, prog, out = {b.movil: b.odo_base for b in db.query(MovilBaseDB)}, {}, []
    for p in db.query(ProgramacionMovilDB): prog.setdefault(p.movil, []).append(p)
    for mv, ls in lect.items():
        m, w, g = _calc(ls, mes0, hasta, c), _calc(ls, d14, hasta, c), _grupo(mv, c)
        paso, ps = c["paso"].get(g), prog.get(mv, [])
        hecho = max([p.hito_km for p in ps if p.estado == "Intervenido"], default=0)
        hito = (int(max(base.get(mv, 0), hecho) // paso) + 1) * paso if paso else 0
        odo = w["odo"] or m["odo"]; dias = (hito - odo) / w["kd"] if hito and w["kd"] > 1 else None
        est = "Sin grupo" if not hito else "Vencido" if odo >= hito else "Próximo" if dias is not None and dias <= hz else "Al día"
        pr = next((p for p in ps if p.estado == "Programado" and p.hito_km == hito), None)
        ints = sorted((p for p in ps if p.estado == "Intervenido"), key=lambda p: p.fecha_intervencion)
        al = [t for t, ok in (("sin dato", m["sd"] >= c["dias_sin_dato"]), ("salto odómetro", m["sal"] > 0), ("baja utilización", m["dd"] > 1 and m["kd"] < c["bajo_uso"])) if ok]
        out.append({"movil": mv, "grupo": g, "km_mes": round(m["km"], 1), "km_dia": round(w["kd"], 1), "odometro": round(odo, 1), "hito": hito,
                    "dias": None if dias is None else round(dias, 1), "estado": est, "programado": pr.fecha_programada if pr else "",
                    "ultima": ints[-1].fecha_intervencion if ints else "", "alertas": ", ".join(al)})
    out.sort(key=lambda r: -r["km_mes"])
    for i, r in enumerate(out): r["pos"] = i + 1
    return hasta, out

@router.get("/programacion")
def programacion(horizonte: int = 15, db: Session = Depends(get_db), user=Depends(get_current_user)):
    r = _flota(db, _cfg(db), horizonte)
    if not r: return {"fecha_corte": "", "resumen": {}, "moviles": [], "agenda": []}
    hasta, mv = r
    ints = db.query(func.count(ProgramacionMovilDB.id)).filter(ProgramacionMovilDB.estado == "Intervenido", ProgramacionMovilDB.fecha_intervencion.like(hasta[:7] + "%")).scalar()
    ag, n = {}, lambda e: sum(1 for x in mv if x["estado"] == e)
    for x in mv:
        if x["programado"]: ag.setdefault(x["programado"], []).append(x["movil"])
    return {"fecha_corte": hasta, "moviles": mv, "agenda": [{"fecha": d, "moviles": sorted(m)} for d, m in sorted(ag.items())],
            "resumen": {"moviles": len(mv), "km_mes": round(sum(x["km_mes"] for x in mv)), "vencidos": n("Vencido"), "proximos": n("Próximo"),
                        "programados": sum(1 for x in mv if x["programado"]), "intervenidos": ints,
                        "faltan": sum(1 for x in mv if x["estado"] in ("Vencido", "Próximo") and not x["programado"])}}

# ── Programar / intervenir ─────────────────────────────────
def _prog(db, mv, hito, fecha, who):
    p = db.query(ProgramacionMovilDB).filter_by(movil=mv, hito_km=hito).first()
    if not fecha:
        if p and p.estado == "Programado": db.delete(p)
        return
    if p is None: p = ProgramacionMovilDB(movil=mv, hito_km=hito); db.add(p)
    p.estado, p.fecha_programada, p.updated_by = "Programado", fecha, who

def _interv(db, mv, hito, fecha, obs, who):
    o = db.query(LecturaKmDB.odo_max).filter(LecturaKmDB.movil == mv, LecturaKmDB.odo_max > 0).order_by(LecturaKmDB.fecha.desc()).first()
    p = db.query(ProgramacionMovilDB).filter_by(movil=mv, hito_km=hito).first() or ProgramacionMovilDB(movil=mv, hito_km=hito)
    db.add(p)
    p.estado, p.fecha_intervencion, p.km_intervencion, p.observacion, p.updated_by = "Intervenido", fecha or date.today().isoformat(), int(o[0]) if o else 0, obs or "", who

class ProgramarIn(BaseModel):
    movil: str
    fecha: Optional[str] = None            # vacío = quitar la programación

class IntervenirIn(BaseModel):
    movil: str
    fecha: Optional[str] = None
    observacion: str = ""

class AutoIn(BaseModel):
    capacidad: int = 10
    desde: Optional[str] = None
    horizonte: int = 15

@router.post("/programar")
def programar(b: ProgramarIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    mv = _norm(b.movil); hito = _pend(db, mv, _cfg(db))
    if not hito: raise HTTPException(404, "Móvil sin kilometraje cargado o sin grupo")
    _prog(db, mv, hito, _fecha(b.fecha), _who(user)); db.commit(); return {"ok": True}

@router.post("/intervenir")
def intervenir(b: IntervenirIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    mv = _norm(b.movil); hito = _pend(db, mv, _cfg(db))
    if not hito: raise HTTPException(404, "Móvil sin kilometraje cargado o sin grupo")
    _interv(db, mv, hito, _fecha(b.fecha), b.observacion, _who(user)); db.commit(); return {"ok": True}

@router.post("/autoprogramar")
def autoprogramar(b: AutoIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    r = _flota(db, _cfg(db), b.horizonte)
    if not r: raise HTTPException(400, "Primero carga el kilometraje")
    d, n, cap = date.fromisoformat(_fecha(b.desde) or date.today().isoformat()), 0, max(b.capacidad, 1)
    cola = sorted((x for x in r[1] if x["estado"] in ("Vencido", "Próximo") and not x["programado"]), key=lambda x: -1e9 if x["dias"] is None else x["dias"])
    for x in cola:                                          # los más urgentes primero, `capacidad` móviles por día
        if n and n % cap == 0: d += timedelta(days=1)
        _prog(db, x["movil"], x["hito"], d.isoformat(), _who(user)); n += 1
    db.commit(); return {"programados": n, "hasta": d.isoformat()}

@router.post("/seguimiento")
def seguimiento(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    """Excel/CSV con columnas Móvil | Estado (Intervenido/Programado) | Fecha | Observación (opcional)."""
    rows, c = _rows(file), _cfg(db)
    h = [str(x or "").lower() for x in (rows[0] if rows else [])]
    ci = lambda k: next((i for i, x in enumerate(h) if k in x), -1)
    im, ie, ifc, io_ = max(ci("móvil"), ci("movil")), ci("estado"), ci("fecha"), ci("obs")
    if im < 0 or ie < 0: raise HTTPException(400, "Columnas requeridas: Móvil, Estado, Fecha (y opcional Observación)")
    ok, bad = 0, []
    for r in rows[1:]:
        get = lambda i: r[i] if 0 <= i < len(r) else ""
        mv, e, f = _norm(get(im)), str(get(ie) or "").lower(), _fecha(get(ifc))
        if not mv: continue
        hito = _pend(db, mv, c); good = False
        if hito and e.startswith("interv"): _interv(db, mv, hito, f, str(get(io_) or ""), _who(user)); good = True
        elif hito and e.startswith("prog") and f: _prog(db, mv, hito, f, _who(user)); good = True
        ok += good
        if not good: bad.append(mv)
    db.commit(); return {"aplicados": ok, "no_aplicados": bad[:30]}
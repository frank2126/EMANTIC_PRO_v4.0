# Programación de móviles a partir del check list de operadores · prefijo /api/checklist
# Carga el Excel del check list (día de por medio), extrae los ítems de mantenimiento,
# detecta móviles con el mismo ítem reportado >= N veces por >= M operadores distintos
# (misma regla de la tabla dinámica manual) y gestiona el seguimiento hasta "Intervenido".
import io, logging, re, time, unicodedata
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel
from sqlalchemy.orm import Session

from database import (CargaHistorialDB, ChecklistItemDB, ChecklistReporteDB, ChecklistSeguimientoDB)
from deps import get_current_admin_user, get_current_user, get_db

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/checklist", tags=["checklist"])

MAX_MB = 10
MIN_REPORTES = 2      # regla: veces que aparece el ítem en el móvil…
MIN_OPERADORES = 2    # …con operadores diferentes


def _n(v: Any) -> str:
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


# Catálogo: nombre canónico + regla de reconocimiento (sobre texto normalizado sin tildes)
CATALOGO = [
    ("Botones de rueda de dirección",            lambda n: "boton" in n and "direccion" in n),
    ("Plataformas deshabilitadas y funcionalidad", lambda n: "plataforma" in n),
    ("Tapa de combustible",                      lambda n: "tapa de combustible" in n),
    ("Puertos USB y condición",                  lambda n: "usb" in n),
    ("Cámara inversa y condición",               lambda n: "camara inversa" in n),
    ("Asiento del operador y estabilidad",       lambda n: "asiento del operador" in n),
    ("Luces internas, timbre",                   lambda n: n.startswith("luces internas") or n == "timbre"),
    ("Estado válvulas despresurización puertas", lambda n: "valvula" in n and "despresuriz" in n),
    ("Ventiladores y defrostadores y funcionalidad", lambda n: "ventilador" in n or "defrost" in n),
    ("Estado de luces",                          lambda n: n.startswith("estado de luces")),
    ("Espejos externos (Ajuste)",                lambda n: "espejos externos" in n),
    ("Rueda de dirección y condición",           lambda n: n.startswith("rueda de direccion")),
]
NOMBRES = [c[0] for c in CATALOGO]


def parse_items(texto: Any) -> List[str]:
    """'A, B, Luces internas, timbre' → ítems del catálogo reconocidos (sin repetir)."""
    if not texto:
        return []
    t = re.sub(r"luces internas\s*,\s*timbre", "luces internas|timbre", str(texto), flags=re.I)
    out: List[str] = []
    for tok in re.split(r",\s*", t):
        n = _n(tok)
        if not n:
            continue
        for nombre, regla in CATALOGO:
            if regla(n):
                if nombre not in out:
                    out.append(nombre)
                break
    return out


def _movil(v: Any) -> str:
    return re.sub(r"\s+", "", str(v or "").upper())

def _cod(v: Any) -> str:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v or "").strip()

def _dt(v: Any) -> Optional[datetime]:
    if isinstance(v, datetime):
        return v.replace(second=0, microsecond=0)
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).replace(second=0, microsecond=0)
        except ValueError:
            pass
    return None

def _iso(d: Optional[datetime]) -> str:
    return d.strftime("%Y-%m-%d %H:%M") if d else ""

def _fecha_corta(v: Optional[str]) -> str:
    m = re.match(r"^\d{4}-\d{2}-\d{2}", str(v or "").strip())
    return m.group(0) if m else ""

def _user(u) -> str:
    return str(u.get("username") or u.get("sub") or "") if isinstance(u, dict) else str(getattr(u, "username", "") or "")


# ── Carga del Excel ────────────────────────────────────────
@router.post("/upload")
def upload_checklist(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    t0 = time.time()
    nombre = file.filename or ""
    if not re.search(r"\.xlsx$", nombre, re.I):
        raise HTTPException(400, "Sube el Excel del check list en formato .xlsx")
    raw = file.file.read()
    if len(raw) > MAX_MB * 1024 * 1024:
        raise HTTPException(413, f"El archivo supera {MAX_MB} MB")
    from openpyxl import load_workbook
    try:
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        filas = list(ws.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(400, "No se pudo leer el archivo Excel")
    if not filas:
        raise HTTPException(400, "El Excel está vacío")

    head = [_n(c) for c in filas[0]]
    def col(*claves):
        for i, h in enumerate(head):
            if any(c in h for c in claves):
                return i
        return -1
    iv, ie, ii, ifch = col("vehiculo", "movil"), col("estado"), col("items de atencion"), col("fecha de creacion", "fecha")
    ic, inom, iln, ir, ino = col("codigo conductor"), col("nombre de conductor", "nombre"), col("lugar"), col("ruta"), col("notas")
    faltan = [n for n, i in (("Vehículo", iv), ("Estado", ie), ("Items de atención", ii), ("Fecha de creación", ifch)) if i < 0]
    if faltan:
        raise HTTPException(400, f"Faltan columnas: {', '.join(faltan)}")

    g = lambda r, i: r[i] if 0 <= i < len(r) else None
    candidatos, sin_novedad, errores = [], 0, 0
    for r in filas[1:]:
        mv = _movil(g(r, iv))
        if not mv:
            continue
        if "necesita" not in _n(g(r, ie)):
            sin_novedad += 1
            continue
        f = _dt(g(r, ifch))
        if not f:
            errores += 1
            continue
        candidatos.append({"movil": mv, "estado": str(g(r, ie) or "")[:50], "fecha": f, "lugar": str(g(r, iln) or "")[:50],
                           "cod": _cod(g(r, ic))[:20], "nom": str(g(r, inom) or "")[:200], "notas": str(g(r, ino) or ""),
                           "ruta": str(g(r, ir) or "")[:30], "texto": str(g(r, ii) or "")})
    if not candidatos:
        total = len(filas) - 1
        _historial(db, user, nombre, total, 0, 0, errores, time.time() - t0)
        return {"reportes_nuevos": 0, "duplicados": 0, "sin_novedad": sin_novedad, "items_relevantes": 0, "moviles": 0, "errores": errores}

    desde, hasta = min(c["fecha"] for c in candidatos), max(c["fecha"] for c in candidatos)
    existentes = {(x.movil, x.fecha_creacion, x.codigo_conductor)
                  for x in db.query(ChecklistReporteDB).filter(ChecklistReporteDB.fecha_creacion.between(desde, hasta))}
    nuevos = dup = n_items = 0
    moviles = set()
    for c in candidatos:
        k = (c["movil"], c["fecha"], c["cod"])
        if k in existentes:
            dup += 1
            continue
        existentes.add(k)
        rep = ChecklistReporteDB(movil=c["movil"], estado=c["estado"], fecha_creacion=c["fecha"], lugar=c["lugar"],
                                 codigo_conductor=c["cod"], nombre_conductor=c["nom"], notas=c["notas"], ruta=c["ruta"], items_texto=c["texto"])
        db.add(rep)
        db.flush()
        for it in parse_items(c["texto"]):
            db.add(ChecklistItemDB(reporte_id=rep.id, movil=c["movil"], item=it, fecha_creacion=c["fecha"],
                                   codigo_conductor=c["cod"], nombre_conductor=c["nom"]))
            n_items += 1
            moviles.add(c["movil"])
        nuevos += 1
    db.commit()
    _historial(db, user, nombre, len(filas) - 1, nuevos, 0, errores, time.time() - t0)
    return {"reportes_nuevos": nuevos, "duplicados": dup, "sin_novedad": sin_novedad, "items_relevantes": n_items,
            "moviles": len(moviles), "desde": _iso(desde), "hasta": _iso(hasta), "errores": errores}

def _historial(db, user, archivo, total, ins, upd, err, seg):
    db.add(CargaHistorialDB(tipo_carga="checklist", usuario_id=_user(user), nombre_archivo=archivo[:300], total_filas=total,
                            filas_insertadas=ins, filas_actualizadas=upd, filas_error=err,
                            estado="exitoso" if err == 0 else ("parcial" if ins else "fallido"),
                            tiempo_procesamiento_seg=str(round(seg, 2))))
    db.commit()


# ── Cálculo de casos ───────────────────────────────────────
def _estado_seg(db: Session):
    corte, abiertos, hist = {}, {}, {}
    for s in db.query(ChecklistSeguimientoDB).order_by(ChecklistSeguimientoDB.id):
        k = (s.movil, s.item)
        if s.estado == "Intervenido":
            hist.setdefault(k, []).append(s)
            if s.corte_fecha and (k not in corte or s.corte_fecha > corte[k]):
                corte[k] = s.corte_fecha
        else:
            abiertos[k] = s
    return corte, abiertos, hist

def _activos(db: Session, corte):
    total, act = {}, {}
    for it in db.query(ChecklistItemDB).order_by(ChecklistItemDB.fecha_creacion):
        k = (it.movil, it.item)
        total[k] = total.get(k, 0) + 1
        c = corte.get(k)
        if c is None or it.fecha_creacion > c:
            act.setdefault(k, []).append(it)
    return total, act

def _seg_dict(s: Optional[ChecklistSeguimientoDB]):
    if not s:
        return {"id": None, "estado": "Pendiente", "fecha_programada": "", "orden_trabajo": "", "observacion": ""}
    return {"id": s.id, "estado": s.estado, "fecha_programada": s.fecha_programada or "",
            "orden_trabajo": s.orden_trabajo or "", "observacion": s.observacion or ""}

def _casos(db: Session, min_rep: int, min_ops: int, incluir_obs: bool):
    corte, abiertos, hist = _estado_seg(db)
    total, act = _activos(db, corte)
    por_movil: Dict[str, List[dict]] = {}
    for k in set(act) | set(abiertos):
        mv, item = k
        rs = act.get(k, [])
        ops = {r.codigo_conductor for r in rs}
        cumple = len(rs) >= min_rep and len(ops) >= min_ops
        if not (cumple or k in abiertos or incluir_obs):
            continue
        por_movil.setdefault(mv, []).append({
            "item": item, "reportes": len(rs), "operadores_distintos": len(ops), "cumple": cumple,
            "primera": _iso(rs[0].fecha_creacion) if rs else "", "ultima": _iso(rs[-1].fecha_creacion) if rs else "",
            "total_historico": total.get(k, 0), "veces_intervenido": len(hist.get(k, [])),
            "reportes_detalle": [{"fecha": _iso(r.fecha_creacion), "codigo": r.codigo_conductor, "operador": r.nombre_conductor} for r in rs],
            "seguimiento": _seg_dict(abiertos.get(k)),
        })
    out = []
    for mv, its in por_movil.items():
        its.sort(key=lambda x: (-int(x["cumple"]), -x["reportes"], x["item"]))
        cumplen = [i for i in its if i["cumple"]]
        out.append({"movil": mv, "items": its, "n_items": len(cumplen), "total_reportes": sum(i["reportes"] for i in its),
                    "cumple": bool(cumplen),
                    "programado": bool(cumplen) and all(i["seguimiento"]["estado"] == "Programado" for i in cumplen)})
    out.sort(key=lambda m: (-int(m["cumple"]), -m["n_items"], -m["total_reportes"], m["movil"]))  # mayor a menor
    return out, hist

@router.get("/programacion")
def programacion(min_reportes: int = Query(MIN_REPORTES, ge=1, le=20), min_operadores: int = Query(MIN_OPERADORES, ge=1, le=20),
                 incluir_observacion: bool = False, db: Session = Depends(get_db), user=Depends(get_current_user)):
    moviles, hist = _casos(db, min_reportes, min_operadores, incluir_observacion)
    hoy = date.today().strftime("%Y-%m")
    interv_mes = sum(1 for v in hist.values() for s in v if (s.fecha_intervencion or "").startswith(hoy))
    a_prog = [m for m in moviles if m["cumple"]]
    ultimo = db.query(ChecklistReporteDB.fecha_creacion).order_by(ChecklistReporteDB.fecha_creacion.desc()).first()
    return {"moviles": moviles, "catalogo": NOMBRES, "ultimo_reporte": _iso(ultimo[0]) if ultimo else "",
            "resumen": {"moviles_a_programar": len(a_prog), "sin_programar": sum(1 for m in a_prog if not m["programado"]),
                        "programados": sum(1 for m in a_prog if m["programado"]), "intervenidos_mes": interv_mes,
                        "reglas": {"min_reportes": min_reportes, "min_operadores": min_operadores}}}

@router.get("/historial/{movil}")
def historial_movil(movil: str, db: Session = Depends(get_db), user=Depends(get_current_user)):
    mv = _movil(movil)
    reps = db.query(ChecklistReporteDB).filter(ChecklistReporteDB.movil == mv).order_by(ChecklistReporteDB.fecha_creacion.desc()).all()
    items = db.query(ChecklistItemDB).filter(ChecklistItemDB.movil == mv).all()
    por_rep: Dict[int, List[str]] = {}
    for i in items:
        por_rep.setdefault(i.reporte_id, []).append(i.item)
    corte, abiertos, hist = _estado_seg(db)
    _, act = _activos(db, corte)
    resumen = []
    for nombre in NOMBRES:
        tot = sum(1 for i in items if i.item == nombre)
        if not tot:
            continue
        k = (mv, nombre)
        resumen.append({"item": nombre, "total": tot, "ciclo_actual": len(act.get(k, [])),
                        "veces_intervenido": len(hist.get(k, [])), "seguimiento": _seg_dict(abiertos.get(k))})
    resumen.sort(key=lambda x: -x["total"])
    ints = db.query(ChecklistSeguimientoDB).filter(ChecklistSeguimientoDB.movil == mv, ChecklistSeguimientoDB.estado == "Intervenido") \
             .order_by(ChecklistSeguimientoDB.fecha_intervencion.desc()).all()
    return {"movil": mv, "total_reportes": len(reps), "por_item": resumen,
            "reportes": [{"fecha": _iso(r.fecha_creacion), "codigo": r.codigo_conductor, "operador": r.nombre_conductor, "ruta": r.ruta,
                          "lugar": r.lugar, "notas": r.notas, "items": por_rep.get(r.id, []), "texto": r.items_texto} for r in reps],
            "intervenciones": [{"id": s.id, "item": s.item, "fecha": s.fecha_intervencion, "orden_trabajo": s.orden_trabajo,
                                "observacion": s.observacion, "por": s.updated_by} for s in ints]}


# ── Seguimiento / intervención ─────────────────────────────
class SeguimientoIn(BaseModel):
    movil: str
    item: Optional[str] = None          # vacío = todos los ítems que cumplen la regla en ese móvil
    fecha_programada: Optional[str] = None
    orden_trabajo: Optional[str] = None
    observacion: Optional[str] = None

class IntervenirIn(BaseModel):
    movil: str
    item: Optional[str] = None          # vacío = todos los ítems pendientes del móvil
    fecha: Optional[str] = None         # por defecto hoy
    orden_trabajo: Optional[str] = None  # opcional
    observacion: Optional[str] = None

class DeshacerIn(BaseModel):
    id: int

def _set(b) -> set:
    return getattr(b, "model_fields_set", None) or getattr(b, "__fields_set__", set())

def _check_item(item: Optional[str]):
    if item and item not in NOMBRES:
        raise HTTPException(400, "Ítem fuera del catálogo")

@router.post("/seguimiento")
def guardar_seguimiento(b: SeguimientoIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    mv = _movil(b.movil); _check_item(b.item)
    if b.item:
        items = [b.item]
    else:
        casos, _ = _casos(db, MIN_REPORTES, MIN_OPERADORES, False)
        items = [i["item"] for m in casos if m["movil"] == mv for i in m["items"] if i["cumple"]]
    if not items:
        raise HTTPException(404, "El móvil no tiene ítems pendientes")
    _, abiertos, _ = _estado_seg(db)
    campos = _set(b)
    for it in items:
        s = abiertos.get((mv, it)) or ChecklistSeguimientoDB(movil=mv, item=it)
        if s.id is None:
            db.add(s)
        if "fecha_programada" in campos:
            s.fecha_programada = _fecha_corta(b.fecha_programada)
        if "orden_trabajo" in campos:
            s.orden_trabajo = (b.orden_trabajo or "").strip()[:50]
        if "observacion" in campos:
            s.observacion = (b.observacion or "").strip()[:300]
        s.estado = "Programado" if s.fecha_programada else "Pendiente"
        s.updated_by = _user(user)
    db.commit()
    return {"ok": True, "actualizados": len(items)}

@router.post("/intervenir")
def intervenir(b: IntervenirIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    mv = _movil(b.movil); _check_item(b.item)
    corte, abiertos, _ = _estado_seg(db)
    _, act = _activos(db, corte)
    claves = [(mv, b.item)] if b.item else sorted({k for k in set(act) | set(abiertos) if k[0] == mv})
    if not claves:
        raise HTTPException(404, "El móvil no tiene ítems pendientes")
    fecha = _fecha_corta(b.fecha) or date.today().isoformat()
    for k in claves:
        rs = act.get(k, [])
        s = abiertos.get(k) or ChecklistSeguimientoDB(movil=k[0], item=k[1])
        if s.id is None:
            db.add(s)
        s.estado = "Intervenido"
        s.fecha_intervencion = fecha
        s.corte_fecha = max((r.fecha_creacion for r in rs), default=datetime.now())   # lo reportado hasta hoy queda atendido
        if b.orden_trabajo is not None and b.orden_trabajo.strip():
            s.orden_trabajo = b.orden_trabajo.strip()[:50]
        if b.observacion is not None and b.observacion.strip():
            s.observacion = b.observacion.strip()[:300]
        s.updated_by = _user(user)
    db.commit()
    return {"ok": True, "intervenidos": len(claves)}

@router.post("/deshacer")
def deshacer(b: DeshacerIn, db: Session = Depends(get_db), user=Depends(get_current_admin_user)):
    """Revierte una intervención marcada por error (los reportes vuelven a contar en el ciclo actual)."""
    s = db.query(ChecklistSeguimientoDB).filter(ChecklistSeguimientoDB.id == b.id, ChecklistSeguimientoDB.estado == "Intervenido").first()
    if not s:
        raise HTTPException(404, "Intervención no encontrada")
    db.delete(s); db.commit()
    return {"ok": True}